"""Render the bilingual-format project report Markdown into a styled DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "Final_ADY201m_Report.md"
OUTPUT = ROOT / "docs" / "Final_ADY201m_Report.docx"


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace(chr(96), "").replace("*", "")
    return text.strip()


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, header=False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(clean_inline(text))
    run.font.name = "Arial"
    run.font.size = Pt(9 if not header else 9.5)
    run.font.bold = header
    run.font.color.rgb = RGBColor(255, 255, 255) if header else RGBColor(42, 52, 66)
    if header:
        shade(cell, "142648")


def add_table(doc: Document, source_lines: list[str]) -> None:
    matrix = []
    for line in source_lines:
        vals = [x.strip() for x in line.strip().strip("|").split("|")]
        if vals and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in vals):
            continue
        matrix.append(vals)
    if not matrix:
        return
    cols = max(len(row) for row in matrix)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    table.autofit = True
    for j in range(cols):
        set_cell_text(table.rows[0].cells[j], matrix[0][j] if j < len(matrix[0]) else "", header=True)
    for row in matrix[1:]:
        cells = table.add_row().cells
        for j in range(cols):
            set_cell_text(cells[j], row[j] if j < len(row) else "")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def set_font(run, *, size=10.5, bold=False, color=RGBColor(42, 52, 66)):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.68)
    section.bottom_margin = Inches(.68)
    section.left_margin = Inches(.72)
    section.right_margin = Inches(.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in [("Title", 23), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(20, 38, 72)
        style.paragraph_format.keep_with_next = True
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.title = "Phân tích giá niêm yết xe đã qua sử dụng tại UAE"
    doc.core_properties.subject = "ADY201m Data Science with Python and SQL"
    doc.core_properties.author = "ADY201m project"

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, block)
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, rel = image_match.groups()
            image_path = (SOURCE.parent / rel).resolve()
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_together = True
                p.add_run().add_picture(str(image_path), width=Inches(6.25))
                cap = doc.add_paragraph(clean_inline(alt))
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(8)
                for run in cap.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(95, 109, 126)
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.add_run(clean_inline(line[2:]))
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(line[2:]))
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(re.sub(r"^\d+\.\s*", "", line)))
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.25)
            r = p.add_run(clean_inline(line[1:].strip()))
            r.italic = True
            i += 1
            continue
        p = doc.add_paragraph()
        p.add_run(clean_inline(line))
        i += 1
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
